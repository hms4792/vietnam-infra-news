#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Vietnam Infrastructure News - Dashboard Updater
Maintains original dashboard format while fixing data loading issues.

Can be run directly: python dashboard_updater.py
"""

import json
import logging
import sys
import os
from datetime import datetime
from pathlib import Path

# Setup paths
SCRIPT_DIR = Path(__file__).parent
PROJECT_ROOT = SCRIPT_DIR.parent
DATA_DIR = PROJECT_ROOT / "data"
OUTPUT_DIR = PROJECT_ROOT / "outputs"
TEMPLATE_DIR = PROJECT_ROOT / "templates"

sys.path.insert(0, str(PROJECT_ROOT))

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(message)s')
logger = logging.getLogger(__name__)

# Excel database path
EXCEL_DB_PATH = DATA_DIR / "database" / "Vietnam_Infra_News_Database_Final.xlsx"

# Template paths
TEMPLATE_HEADER = TEMPLATE_DIR / "dashboard_template_header.html"
TEMPLATE_FOOTER = TEMPLATE_DIR / "dashboard_template_footer.html"


def load_articles_from_excel():
    """Load ALL articles from Excel database"""
    try:
        import openpyxl
    except ImportError:
        logger.error("openpyxl not installed")
        return []
    
    if not EXCEL_DB_PATH.exists():
        logger.warning(f"Excel database not found: {EXCEL_DB_PATH}")
        return []
    
    logger.info(f"Loading from: {EXCEL_DB_PATH}")
    
    try:
        wb = openpyxl.load_workbook(EXCEL_DB_PATH, read_only=True, data_only=True)
        ws = wb.active
        
        # Get headers
        headers = [cell.value for cell in ws[1]]
        col_map = {str(h).strip(): i for i, h in enumerate(headers) if h}
        
        logger.info(f"Headers: {list(col_map.keys())}")
        
        articles = []
        for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not any(row):
                continue
            
            # Parse date
            date_val = row[col_map.get("Date", 4)] if "Date" in col_map else None
            if date_val:
                if hasattr(date_val, 'strftime'):
                    date_str = date_val.strftime("%Y-%m-%d")
                else:
                    date_str = str(date_val)[:10]
            else:
                date_str = ""
            
            # Get values
            area = row[col_map.get("Area", 0)] or "Environment"
            sector = row[col_map.get("Business Sector", 1)] or "Unknown"
            province = row[col_map.get("Province", 2)] or "Vietnam"
            title = str(row[col_map.get("News Tittle", 3)] or "")
            source = row[col_map.get("Source", 5)] or ""
            url = row[col_map.get("Link", 6)] or ""
            summary = str(row[col_map.get("Short summary", 7)] or "")
            
            if not title and not url:
                continue
            
            articles.append({
                "id": len(articles) + 1,
                "date": date_str,
                "area": area,
                "sector": sector,
                "province": province,
                "source": source,
                "title": title,
                "summary": summary,
                "url": url
            })
        
        wb.close()
        
        # Statistics
        from collections import Counter
        year_counts = Counter(a["date"][:4] for a in articles if a["date"])
        sector_counts = Counter(a["sector"] for a in articles)
        
        logger.info(f"Loaded {len(articles)} articles")
        logger.info(f"Years: {dict(sorted(year_counts.items()))}")
        logger.info(f"Top sectors: {dict(sector_counts.most_common(5))}")
        
        return articles
        
    except Exception as e:
        logger.error(f"Error loading Excel: {e}")
        import traceback
        traceback.print_exc()
        return []


def convert_to_multilingual_format(articles):
    """
    Convert articles to multilingual format for BACKEND_DATA.
    Original format:
    {
        "id": 1,
        "date": "2026-01-23",
        "area": "Environment",
        "sector": "Waste Water",
        "province": "Vietnam",
        "source": "VnExpress",
        "title": {"ko": "...", "en": "...", "vi": "..."},
        "summary": {"ko": "...", "en": "...", "vi": "..."},
        "url": "https://..."
    }
    """
    result = []
    
    # 섹터 한국어 번역
    SECTOR_KO = {
        "Waste Water": "폐수처리",
        "Solid Waste": "고형폐기물",
        "Water Supply/Drainage": "상수도/배수",
        "Power": "발전/전력",
        "Oil & Gas": "석유/가스",
        "Industrial Parks": "산업단지",
        "Smart City": "스마트시티",
        "Transport": "교통인프라",
        "Climate Change": "기후변화"
    }
    
    # 섹터 베트남어 번역
    SECTOR_VI = {
        "Waste Water": "Xử lý nước thải",
        "Solid Waste": "Chất thải rắn",
        "Water Supply/Drainage": "Cấp thoát nước",
        "Power": "Điện năng",
        "Oil & Gas": "Dầu khí",
        "Industrial Parks": "Khu công nghiệp",
        "Smart City": "Thành phố thông minh",
        "Transport": "Giao thông",
        "Climate Change": "Biến đổi khí hậu"
    }
    
    for article in articles:
        title = article.get("title", "")
        summary = article.get("summary", "")
        sector = article.get("sector", "Infrastructure")
        province = article.get("province", "Vietnam")
        
        sector_ko = SECTOR_KO.get(sector, sector)
        sector_vi = SECTOR_VI.get(sector, sector)
        
        # 제목은 원본 유지 (번역 없음)
        title_ko = title
        title_en = title
        title_vi = title
        
        # 기존 요약이 있는지 확인
        has_existing_summary = summary and len(summary.strip()) > 20
        
        # 언어 감지 (베트남어 문자 포함 여부)
        has_vietnamese_chars = any(ord(c) > 127 for c in title)
        
        # 요약 생성
        if has_existing_summary:
            # 기존 요약이 있으면 활용
            base_summary = summary[:300]
            
            # 기존 요약의 언어 감지
            if "project in Vietnam" in summary or "Vietnam" in summary[:50]:
                # 영어 요약
                summary_en = base_summary
                summary_ko = f"[{sector_ko}] {province} - {title[:80]}"
                summary_vi = f"[{sector_vi}] {province} - {title[:80]}"
            else:
                # 베트남어 또는 기타 요약
                summary_vi = base_summary
                summary_en = f"[{sector}] {province} - {title[:80]}"
                summary_ko = f"[{sector_ko}] {province} - {title[:80]}"
        else:
            # 기존 요약이 없으면 템플릿 생성
            if has_vietnamese_chars:
                summary_vi = f"[{sector_vi}] {province}: {title[:150]}"
                summary_en = f"[{sector}] {province}: {title[:150]}"
                summary_ko = f"[{sector_ko}] {province}: {title[:150]}"
            else:
                summary_en = f"[{sector}] {province}: {title[:150]}"
                summary_ko = f"[{sector_ko}] {province}: {title[:150]}"
                summary_vi = f"[{sector_vi}] {province}: {title[:150]}"
        
        result.append({
            "id": article.get("id", len(result) + 1),
            "date": article.get("date", ""),
            "area": article.get("area", "Environment"),
            "sector": sector,
            "province": province,
            "source": article.get("source", ""),
            "title": {
                "ko": title_ko,
                "en": title_en,
                "vi": title_vi
            },
            "summary": {
                "ko": summary_ko,
                "en": summary_en,
                "vi": summary_vi
            },
            "url": article.get("url", "")
        })
    
    return result


def generate_dashboard(articles):
    """Generate dashboard HTML using original template format"""
    
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    
    # Convert to multilingual format
    backend_data = convert_to_multilingual_format(articles)
    
    # Sort by date (newest first)
    backend_data.sort(key=lambda x: x.get("date", ""), reverse=True)
    
    # Generate JSON for BACKEND_DATA
    backend_json = json.dumps(backend_data, ensure_ascii=False, indent=2)
    
    # Check if template exists
    if TEMPLATE_HEADER.exists() and TEMPLATE_FOOTER.exists():
        logger.info("Using template files")
        
        with open(TEMPLATE_HEADER, 'r', encoding='utf-8') as f:
            header = f.read()
        
        with open(TEMPLATE_FOOTER, 'r', encoding='utf-8') as f:
            footer = f.read()
        
        # Combine: header + BACKEND_DATA + footer
        html = header + f"\nconst BACKEND_DATA = {backend_json};\n" + footer
        
    else:
        logger.info("Template not found, generating minimal dashboard")
        html = generate_minimal_dashboard(backend_data, backend_json)
    
    # Save files
    index_path = OUTPUT_DIR / "index.html"
    dashboard_path = OUTPUT_DIR / "vietnam_dashboard.html"
    
    with open(index_path, 'w', encoding='utf-8') as f:
        f.write(html)
    logger.info(f"✓ Saved: {index_path} ({index_path.stat().st_size:,} bytes)")
    
    with open(dashboard_path, 'w', encoding='utf-8') as f:
        f.write(html)
    logger.info(f"✓ Saved: {dashboard_path}")
    
    # Also save JSON data
    json_path = OUTPUT_DIR / f"news_data_{datetime.now().strftime('%Y%m%d')}.json"
    with open(json_path, 'w', encoding='utf-8') as f:
        json.dump({
            "generated_at": datetime.now().isoformat(),
            "total": len(backend_data),
            "new_articles": sum(1 for a in backend_data if a.get("date") == datetime.now().strftime("%Y-%m-%d"))
        }, f, indent=2)
    
    return str(index_path)


def generate_minimal_dashboard(articles, backend_json):
    """Generate minimal dashboard when template is not available"""
    
    today = datetime.now().strftime("%Y-%m-%d")
    total = len(articles)
    today_count = sum(1 for a in articles if a.get("date") == today)
    
    html = f'''<!DOCTYPE html>
<html lang="ko">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Vietnam Infrastructure News Dashboard</title>
    <link href="https://fonts.googleapis.com/css2?family=Noto+Sans+KR:wght@400;500;600;700&display=swap" rel="stylesheet">
    <script src="https://cdn.tailwindcss.com"></script>
    <style>
        * {{ font-family: 'Noto Sans KR', sans-serif; }}
        .lang-btn.active {{ background: linear-gradient(135deg, #0F766E, #14B8A6); color: white; }}
        .filter-btn.active {{ background: #0F766E; color: white; }}
        .sector-env {{ background: linear-gradient(135deg, #059669, #10B981); }}
        .sector-energy {{ background: linear-gradient(135deg, #D97706, #F59E0B); }}
        .sector-urban {{ background: linear-gradient(135deg, #7C3AED, #8B5CF6); }}
    </style>
</head>
<body class="bg-gradient-to-br from-slate-50 to-teal-50/30 min-h-screen">

<header class="sticky top-0 z-50 bg-white/90 backdrop-blur border-b border-slate-200">
    <div class="max-w-7xl mx-auto px-4 py-3">
        <div class="flex items-center justify-between">
            <div class="flex items-center gap-3">
                <div class="w-10 h-10 bg-gradient-to-br from-teal-600 to-emerald-500 rounded-lg flex items-center justify-center text-white text-xl">🏗️</div>
                <div>
                    <h1 class="text-lg font-bold text-slate-800">Vietnam Infra News</h1>
                    <p class="text-xs text-slate-500">Updated: {datetime.now().strftime("%Y-%m-%d %H:%M")}</p>
                </div>
            </div>
            <div class="flex gap-1 bg-slate-100 rounded-lg p-0.5">
                <button onclick="setLang('ko')" class="lang-btn active px-3 py-1.5 rounded text-sm font-medium" data-lang="ko">🇰🇷 한국어</button>
                <button onclick="setLang('en')" class="lang-btn px-3 py-1.5 rounded text-sm font-medium text-slate-600" data-lang="en">🇺🇸 EN</button>
                <button onclick="setLang('vi')" class="lang-btn px-3 py-1.5 rounded text-sm font-medium text-slate-600" data-lang="vi">🇻🇳 VI</button>
            </div>
        </div>
    </div>
</header>

<main class="max-w-7xl mx-auto px-4 py-6">
    <!-- KPI Cards -->
    <div class="grid grid-cols-2 md:grid-cols-4 gap-4 mb-6">
        <div class="bg-white rounded-xl p-4 shadow-sm border">
            <div class="text-2xl font-bold text-teal-600">{today_count}</div>
            <div class="text-sm text-slate-500">오늘</div>
        </div>
        <div class="bg-white rounded-xl p-4 shadow-sm border">
            <div class="text-2xl font-bold text-blue-600">{total:,}</div>
            <div class="text-sm text-slate-500">전체 DB</div>
        </div>
    </div>
    
    <!-- News List -->
    <div class="bg-white rounded-xl shadow-lg overflow-hidden">
        <div class="bg-gradient-to-r from-teal-700 to-emerald-600 px-4 py-3">
            <span class="text-white font-bold">📰 뉴스 목록</span>
            <span class="text-teal-100 text-sm ml-2" id="news-count">{total}건</span>
        </div>
        <div id="news-list" class="max-h-[600px] overflow-y-auto"></div>
    </div>
</main>

<script>
const BACKEND_DATA = {backend_json};

let lang = 'ko';

function setLang(l) {{
    lang = l;
    document.querySelectorAll('.lang-btn').forEach(btn => {{
        btn.classList.toggle('active', btn.dataset.lang === l);
    }});
    renderNews();
}}

function getText(obj) {{
    if (!obj) return '';
    if (typeof obj === 'string') return obj;
    return obj[lang] || obj.en || obj.ko || obj.vi || '';
}}

function getSectorColor(area) {{
    if (area === 'Environment') return 'sector-env';
    if (area === 'Energy Develop.') return 'sector-energy';
    return 'sector-urban';
}}

function renderNews() {{
    const container = document.getElementById('news-list');
    container.innerHTML = BACKEND_DATA.slice(0, 100).map(n => `
        <div class="px-4 py-3 border-b hover:bg-slate-50">
            <div class="flex items-center gap-2 mb-1">
                <span class="px-2 py-0.5 rounded text-xs font-semibold text-white ${{getSectorColor(n.area)}}">${{n.sector}}</span>
                <span class="text-xs text-slate-500">${{n.date}}</span>
                <span class="text-xs text-slate-400">${{n.source}}</span>
            </div>
            <h3 class="font-medium text-slate-800">${{getText(n.title)}}</h3>
            <p class="text-sm text-slate-600 mt-1">${{getText(n.summary).slice(0, 150)}}...</p>
            <a href="${{n.url}}" target="_blank" class="text-xs text-teal-600 hover:underline mt-1 inline-block">원문보기 →</a>
        </div>
    `).join('');
}}

renderNews();
</script>
</body>
</html>'''
    
    return html


def main():
    """Main function - generates dashboard from Excel"""
    print("=" * 60)
    print("DASHBOARD GENERATOR (Original Format)")
    print("=" * 60)
    print(f"Time: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    print(f"Excel: {EXCEL_DB_PATH}")
    print(f"Template header: {TEMPLATE_HEADER}")
    print(f"Template footer: {TEMPLATE_FOOTER}")
    
    # Ensure directories
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    
    # Load articles
    articles = load_articles_from_excel()
    
    if not articles:
        print("WARNING: No articles loaded from Excel!")
        print("Generating empty dashboard...")
    
    # Generate dashboard
    print(f"\nGenerating dashboard with {len(articles)} articles...")
    result = generate_dashboard(articles)
    
    # Verify
    index_path = OUTPUT_DIR / "index.html"
    if index_path.exists():
        print(f"\n✓ Dashboard created: {index_path}")
        print(f"  Size: {index_path.stat().st_size:,} bytes")
    else:
        print("\n✗ Dashboard generation failed!")
        return
    
    print("\nDone!")


if __name__ == "__main__":
    main()
